"""Read an .xlsx file into `RawSheet`s. The only module that touches openpyxl when reading.

Excel stores a colour as RGB, as a theme slot plus a tint, or as a legacy palette index;
all three are resolved to plain RGB here so the rest of the code only sees `Rgb`.
"""

import colorsys
import logging
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell as ExcelCell
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.formatting.rule import Rule
from openpyxl.styles.colors import COLOR_INDEX, Color
from openpyxl.styles.fills import GradientFill, PatternFill
from openpyxl.utils.cell import column_index_from_string, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from validation_checks.config import (
    EXCEL_MAX_COLUMNS,
    EXCEL_MAX_ROWS,
    THEME_NAMESPACE,
    THEME_SLOTS,
    UNSET_RGB,
)
from validation_checks.models import (
    Cell,
    CellPos,
    CellRange,
    CellValue,
    CfRule,
    DataLoadError,
    RawSheet,
    Rgb,
    Style,
    TextRun,
)

logger = logging.getLogger(__name__)


@dataclass(frozen=True, slots=True, kw_only=True)
class _Palette:
    theme: tuple[Rgb | None, ...]
    indexed: tuple[str, ...]


def read_workbook(path: Path) -> tuple[RawSheet, ...]:
    try:
        workbook = load_workbook(path, data_only=True, rich_text=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as error:
        raise DataLoadError(f"Cannot open workbook: {error}") from error
    try:
        custom_palette = getattr(workbook, "_colors", None)  # workbook-specific legacy palette
        palette = _Palette(
            theme=_theme_colours(workbook.loaded_theme),
            indexed=tuple(custom_palette if isinstance(custom_palette, list) else COLOR_INDEX),
        )
        return tuple(
            _read_sheet(sheet, palette)
            for sheet in workbook.worksheets
            if isinstance(sheet, Worksheet)
        )
    finally:
        workbook.close()


# --------------------------------------------------------------------------- sheets and cells


def _read_sheet(sheet: Worksheet, palette: _Palette) -> RawSheet:
    cells: dict[CellPos, Cell] = {}
    for row in sheet.iter_rows():
        for excel_cell in row:
            if not isinstance(excel_cell, MergedCell):
                cell = _read_cell(excel_cell, palette)
                if cell is not None:
                    cells[CellPos(excel_cell.row, excel_cell.column)] = cell
    rules = tuple(
        _read_rule(rule, tuple(_range(str(ref)) for ref in formatting.sqref.ranges), palette)
        for formatting in sheet.conditional_formatting
        for rule in formatting.rules
    )
    hidden_columns: set[int] = set()
    for key, dimension in sheet.column_dimensions.items():
        if dimension.hidden:
            start = dimension.min or column_index_from_string(key)
            hidden_columns.update(range(start, (dimension.max or start) + 1))
    return RawSheet(
        name=sheet.title,
        cells=cells,
        max_row=sheet.max_row,
        max_column=sheet.max_column,
        cf_rules=rules,
        hidden_rows=frozenset(i for i, dim in sheet.row_dimensions.items() if dim.hidden),
        hidden_columns=frozenset(hidden_columns),
    )


def _read_cell(excel_cell: ExcelCell, palette: _Palette) -> Cell | None:
    # cell.fill is a StyleProxy, which fails isinstance checks; copy() unwraps the real fill.
    fill = _fill_colour(copy(excel_cell.fill), palette, conditional=False)
    font = _colour(excel_cell.font.color, palette) if excel_cell.font else None
    strike = bool(excel_cell.font.strike) if excel_cell.font else False
    value: object = excel_cell.value
    if value is None and fill is None and not strike:
        return None
    runs: tuple[TextRun, ...] = ()
    if isinstance(value, CellRichText):
        runs = tuple(_run(part, font, strike, palette) for part in value)
    return Cell(value=_plain(value), style=Style(fill=fill, font=font, strike=strike), runs=runs)


def _run(part: str | TextBlock, font: Rgb | None, strike: bool, palette: _Palette) -> TextRun:
    """A rich-text run; anything the run does not set is inherited from the cell font."""
    if isinstance(part, str):
        return TextRun(text=part, font=font, strike=strike)
    inline = part.font
    return TextRun(
        text=part.text,
        font=_colour(inline.color, palette) if inline.color is not None else font,
        strike=strike if inline.strike is None else bool(inline.strike),
    )


def _plain(value: object) -> CellValue:
    match value:
        case None | str() | int() | float():
            return value
        case datetime() | date() | time():
            return value.isoformat()
        case _:
            return str(value)


def _read_rule(rule: Rule, ranges: tuple[CellRange, ...], palette: _Palette) -> CfRule:
    style = Style()
    if (dxf := rule.dxf) is not None:
        font = dxf.font
        style = Style(
            fill=_fill_colour(dxf.fill, palette, conditional=True) if dxf.fill else None,
            font=_colour(font.color, palette) if font is not None and font.color else None,
            strike=None if font is None or font.strike is None else bool(font.strike),
        )
    return CfRule(
        kind=str(rule.type),
        operator=None if rule.operator is None else str(rule.operator),
        operands=tuple(str(formula) for formula in (rule.formula or ())),
        text=None if rule.text is None else str(rule.text),
        style=style,
        priority=int(rule.priority or 0),
        stop_if_true=bool(rule.stopIfTrue),
        ranges=ranges,
    )


def _range(reference: str) -> CellRange:
    min_column, min_row, max_column, max_row = range_boundaries(reference)
    return CellRange(
        min_row=min_row or 1,
        min_column=min_column or 1,
        max_row=max_row or EXCEL_MAX_ROWS,
        max_column=max_column or EXCEL_MAX_COLUMNS,
    )


# --------------------------------------------------------------------------- colours


def _fill_colour(fill: object, palette: _Palette, *, conditional: bool) -> Rgb | None:
    if isinstance(fill, PatternFill):
        if conditional:  # conditional fills keep the visible colour in bgColor (or fgColor)
            for colour in (fill.bgColor, fill.fgColor):
                if colour is not None and not (colour.type == "rgb" and colour.rgb == UNSET_RGB):
                    return _colour(colour, palette)
            return None
        return None if fill.fill_type in {None, "none"} else _colour(fill.fgColor, palette)
    if isinstance(fill, GradientFill) and fill.stop:
        return _colour(fill.stop[0].color, palette)
    return None


def _colour(colour: Color | None, palette: _Palette) -> Rgb | None:
    if colour is None:
        return None
    match colour.type:
        case "rgb":
            return Rgb.from_hex(colour.rgb) if isinstance(colour.rgb, str) else None
        case "theme":
            index = int(colour.theme)
            base = palette.theme[index] if 0 <= index < len(palette.theme) else None
            return apply_tint(base, float(colour.tint or 0.0)) if base else None
        case "indexed":
            index = int(colour.indexed)
            in_palette = 0 <= index < len(palette.indexed)
            return Rgb.from_hex(palette.indexed[index]) if in_palette else None
        case _:
            return None


def apply_tint(colour: Rgb, tint: float) -> Rgb:
    """Excel theme tint (-1..1) moves HLS lightness towards black (<0) or white (>0)."""
    hue, lightness, saturation = colorsys.rgb_to_hls(
        colour.red / 255, colour.green / 255, colour.blue / 255
    )
    lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
    red, green, blue = colorsys.hls_to_rgb(hue, lightness, saturation)
    return Rgb(round(red * 255), round(green * 255), round(blue * 255))


def _theme_colours(theme_xml: bytes | str | None) -> tuple[Rgb | None, ...]:
    """Colours indexed by Excel theme number, read from the workbook's theme XML."""
    if not theme_xml:
        return ()
    try:
        scheme = ElementTree.fromstring(theme_xml).find(".//a:clrScheme", THEME_NAMESPACE)
    except ElementTree.ParseError:
        logger.warning("Workbook theme is unreadable; theme colours will not be resolved")
        return ()
    if scheme is None:
        return ()
    colours: list[Rgb | None] = []
    for slot_name in THEME_SLOTS:
        slot = scheme.find(f"a:{slot_name}", THEME_NAMESPACE)
        children = list(slot) if slot is not None else []
        # <a:srgbClr val="C0504D"/> or <a:sysClr val="windowText" lastClr="000000"/>
        value = next((child.get("lastClr") or child.get("val") or "" for child in children), "")
        colours.append(Rgb.from_hex(value))
    return tuple(colours)
