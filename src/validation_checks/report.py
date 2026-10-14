"""Write the results workbook: a fresh, timestamped file per run with fixed sheets.

Every sheet is always written (headers only when empty), so "no issues" is never confused with
"check did not run". Sheet names and columns come from `REPORT_COLUMNS` in config.py.
"""

from collections.abc import Iterable, Sequence
from datetime import datetime
from pathlib import Path
from typing import Self, TypeAlias

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from validation_checks.config import MAX_COLUMN_WIDTH, REPORT_COLUMNS, REPORT_FILE
from validation_checks.models import FileReport, Severity

Row: TypeAlias = Sequence[str | int | float | None]


class ReportBuilder:
    """Fluent builder: add sheets in order, then save."""

    def __init__(self) -> None:
        self._workbook = Workbook()
        self._workbook.remove(self._workbook.worksheets[0])

    def add_sheet(self, title: str, rows: Iterable[Row]) -> Self:
        headers = REPORT_COLUMNS[title]
        sheet = self._workbook.create_sheet(title)
        sheet.append(list(headers))
        widths = [len(header) for header in headers]
        for row in rows:
            sheet.append([str(v) if isinstance(v, str) else v for v in row])  # StrEnum -> str
            widths = [max(w, len(str(v or ""))) for w, v in zip(widths, row, strict=True)]
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                width + 2, MAX_COLUMN_WIDTH
            )
        return self

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self._workbook.save(path)
        return path


def write_report(reports: Sequence[FileReport], output_dir: Path, started: datetime) -> Path:
    return (
        ReportBuilder()
        .add_sheet("Summary", _summary(reports))
        .add_sheet("Issues", _issues(reports))
        .add_sheet("LOB Matching", _matches(reports))
        .add_sheet("Exclusions", _exclusions(reports))
        .add_sheet("File Errors", _file_errors(reports))
        .save(output_dir / REPORT_FILE.format(started))
    )


def _summary(reports: Sequence[FileReport]) -> Iterable[Row]:
    for r in reports:
        yield (r.file, r.country, r.status, r.count(Severity.ERROR), r.count(Severity.WARNING),
               len(r.exclusions), r.review_count, len(r.file_errors))  # fmt: skip


def _issues(reports: Sequence[FileReport]) -> Iterable[Row]:
    for r in reports:
        for i in r.issues:
            yield (r.file, r.country, i.sheet, i.row, i.column, i.severity, i.check,
                   i.message, i.value)  # fmt: skip


def _matches(reports: Sequence[FileReport]) -> Iterable[Row]:
    for r in reports:
        for m in r.matches:
            candidate = m.hit.candidate.display if m.hit.candidate else ""
            yield (r.file, r.country, m.direction, m.name.display, candidate, m.hit.method,
                   m.hit.score, m.hit.status, ", ".join(m.name.sources))  # fmt: skip


def _exclusions(reports: Sequence[FileReport]) -> Iterable[Row]:
    for r in reports:
        for e in r.exclusions:
            yield (r.file, e.sheet, e.row, e.column, e.value, e.scope, e.reason, e.colour)


def _file_errors(reports: Sequence[FileReport]) -> Iterable[Row]:
    for r in reports:
        for e in r.file_errors:
            yield (r.file, e.sheet, e.message)
