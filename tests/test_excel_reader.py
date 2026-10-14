"""The reader must resolve every colour encoding Excel uses and keep CF rules intact."""

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.worksheet import Worksheet

from tests.factories import RED_HEX, SETTINGS, pos, red_font, solid
from validation_checks.excel_reader import apply_tint, read_workbook
from validation_checks.exclusions import is_red
from validation_checks.models import DataLoadError, RawSheet, Rgb


def round_trip(tmp_path: Path, workbook: Workbook) -> RawSheet:
    workbook.save(tmp_path / "book.xlsx")
    return read_workbook(tmp_path / "book.xlsx")[0]


def new_sheet() -> tuple[Workbook, Worksheet]:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    return workbook, sheet


def test_rgb_fill_and_font(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    sheet["A1"], sheet["A2"] = "filled", "font"
    sheet["A1"].fill = solid(RED_HEX)
    sheet["A2"].font = red_font()

    raw = round_trip(tmp_path, workbook)

    assert is_red(raw.cell(pos("A1")).style.fill, SETTINGS)
    assert is_red(raw.cell(pos("A2")).style.font, SETTINGS)


@pytest.mark.parametrize(
    ("theme", "tint", "red"), [(5, 0.4, True), (5, 0.0, True), (4, 0.0, False)]
)
def test_theme_colours_resolve_through_workbook_theme(
    tmp_path: Path, theme: int, tint: float, red: bool
) -> None:
    workbook, sheet = new_sheet()
    sheet["A1"] = "themed"
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=theme, tint=tint))

    assert is_red(round_trip(tmp_path, workbook).cell(pos("A1")).style.fill, SETTINGS) is red


def test_indexed_palette_colour(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    sheet["A1"] = "indexed"
    sheet["A1"].font = Font(color=Color(indexed=10))

    assert is_red(round_trip(tmp_path, workbook).cell(pos("A1")).style.font, SETTINGS)


def test_strikethrough(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    sheet["A1"] = "gone"
    sheet["A1"].font = Font(strike=True)

    assert round_trip(tmp_path, workbook).cell(pos("A1")).style.strike is True


def test_rich_text_runs_inherit_cell_font(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    sheet["A1"] = CellRichText("Carrier ", TextBlock(InlineFont(color=Color(RED_HEX)), "A"))
    sheet["A1"].font = Font(strike=True)

    cell = round_trip(tmp_path, workbook).cell(pos("A1"))

    assert cell.text == "Carrier A"
    assert [run.text for run in cell.runs] == ["Carrier ", "A"]
    assert all(run.strike for run in cell.runs)
    assert is_red(cell.runs[1].font, SETTINGS)


def test_conditional_format_rules_are_read(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    fill = PatternFill(fill_type="solid", bgColor=RED_HEX, fgColor=RED_HEX)
    cell_is = Rule(
        type="cellIs", operator="equal", formula=['"X"'], dxf=DifferentialStyle(fill=fill)
    )
    expression = Rule(type="expression", formula=["$A1=1"], dxf=DifferentialStyle(font=red_font()))
    sheet.conditional_formatting.add("A1:A10", cell_is)
    sheet.conditional_formatting.add("B1:B10", expression)

    rules = {rule.kind: rule for rule in round_trip(tmp_path, workbook).cf_rules}

    assert rules["cellIs"].operands == ('"X"',)
    assert is_red(rules["cellIs"].style.fill, SETTINGS)
    assert is_red(rules["expression"].style.font, SETTINGS)


def test_hidden_rows_and_columns(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    sheet["C3"] = "x"
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True

    raw = round_trip(tmp_path, workbook)

    assert (2 in raw.hidden_rows, 2 in raw.hidden_columns) == (True, True)


def test_whole_numbers_read_without_decimals(tmp_path: Path) -> None:
    workbook, sheet = new_sheet()
    sheet["A1"] = 10007481

    assert round_trip(tmp_path, workbook).cell(pos("A1")).text == "10007481"


def test_non_workbook_raises_data_load_error(tmp_path: Path) -> None:
    (tmp_path / "broken.xlsx").write_text("not a zip", encoding="utf-8")

    with pytest.raises(DataLoadError):
        read_workbook(tmp_path / "broken.xlsx")


@pytest.mark.parametrize(("tint", "expected"), [(1.0, Rgb(255, 255, 255)), (-1.0, Rgb(0, 0, 0))])
def test_tint_moves_towards_white_or_black(tint: float, expected: Rgb) -> None:
    assert apply_tint(Rgb(192, 0, 0), tint) == expected
