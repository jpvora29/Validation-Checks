"""Test inputs: in-memory sheets for logic tests, real workbooks for reader and end-to-end tests."""

import re
from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from validation_checks.config import Settings
from validation_checks.models import (
    Cell,
    CellPos,
    CellValue,
    CfRule,
    RawSheet,
    Rgb,
    Style,
    TextRun,
    column_index,
)

SETTINGS = Settings()
RED = Rgb(255, 0, 0)
GREEN = Rgb(0, 176, 80)
RED_HEX = "FFFF0000"
GREEN_HEX = "FF00B050"


def pos(coordinate: str) -> CellPos:
    match = re.match(r"^([A-Z]+)(\d+)$", coordinate)
    assert match is not None, coordinate
    return CellPos(int(match[2]), column_index(match[1]))


def make_sheet(
    rows: Sequence[Sequence[CellValue]],
    *,
    name: str = "Sheet",
    styles: Mapping[str, Style] | None = None,
    runs: Mapping[str, tuple[TextRun, ...]] | None = None,
    cf_rules: tuple[CfRule, ...] = (),
    hidden_rows: frozenset[int] = frozenset(),
) -> RawSheet:
    """A RawSheet from row lists (row 1 first); `styles` / `runs` are keyed like "B3"."""
    cells = {
        CellPos(r, c): Cell(value=value)
        for r, row in enumerate(rows, start=1)
        for c, value in enumerate(row, start=1)
    }
    for coordinate, style in (styles or {}).items():
        cells[pos(coordinate)] = Cell(value=cells.get(pos(coordinate), Cell()).value, style=style)
    for coordinate, cell_runs in (runs or {}).items():
        existing = cells.get(pos(coordinate), Cell())
        cells[pos(coordinate)] = Cell(value=existing.value, style=existing.style, runs=cell_runs)
    return RawSheet(
        name=name,
        cells=cells,
        max_row=len(rows),
        max_column=max((len(row) for row in rows), default=0),
        cf_rules=cf_rules,
        hidden_rows=hidden_rows,
    )


# Template-shaped rows: legend on row 1, header on row 7, data from row 8.
LEGEND_ROWS: tuple[tuple[CellValue, ...], ...] = (
    ("Title", None, "EXCLUDE", "INCLUDE"),
    ("Purpose",),
    (),
    ("Instructions",),
    (),
    (),
)
LEGEND_STYLES: Mapping[str, Style] = {"C1": Style(fill=RED), "D1": Style(fill=GREEN)}
PEERS_HEADER = (
    "Region", "Country", "Carrier Group", "Carrier Name (as it would appear in Survey)",
    "Carrier Entity", "ICA carrier", "Peer 1", "Peer  2", "Peer 3", "Peer 4", "Peer 5", "Peer 6",
)  # fmt: skip
LOB_HEADER = (
    "Region",
    "Country",
    "Carrier",
    "ICA carrier",
    "CE/CM",
    "Claims",
    "Property",
    "Surety",
)
PANEL_HEADER = ("Employee ID", "First Name", "Last Name", "Email", "Country", "Job Title")


def peers_row(carrier: str, peers: Sequence[str], *, group: str = "GROUP") -> tuple[CellValue, ...]:
    return ("LAC", "Argentina", group, carrier, "Entity", "X", *peers)


def lob_row(carrier: str, practices: Sequence[str | None]) -> tuple[CellValue, ...]:
    return ("LAC", "Argentina", carrier, "X", *practices)


def panel_row(employee_id: int, email: str, *, last: str = "Diaz") -> tuple[CellValue, ...]:
    return (employee_id, "Ana", last, email, "Argentina", "Associate")


def survey_workbook(
    *,
    peers: Sequence[Sequence[CellValue]],
    lob: Sequence[Sequence[CellValue]],
    panel: Sequence[Sequence[CellValue]],
) -> Workbook:
    """A real workbook shaped like the template, with the red/green legend on row 1."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for title, header, rows in (
        ("1.Country Peers", PEERS_HEADER, peers),
        ("2. Country LOB MAPPING.", LOB_HEADER, lob),
        ("3. User Panel", PANEL_HEADER, panel),
    ):
        sheet = workbook.create_sheet(title)
        for row in (*LEGEND_ROWS, header, *rows):
            sheet.append(list(row))
        sheet["C1"].fill = solid(RED_HEX)
        sheet["D1"].fill = solid(GREEN_HEX)
    return workbook


def solid(hex_colour: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_colour)


def red_font() -> Font:
    return Font(color=RED_HEX)


def save(workbook: Workbook, path: Path) -> Path:
    workbook.save(path)
    return path
