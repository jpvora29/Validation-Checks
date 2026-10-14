"""Whole runs through the real Excel reader, the shipped sample, and the CLI."""

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.factories import (
    SETTINGS,
    lob_row,
    panel_row,
    peers_row,
    red_font,
    save,
    solid,
    survey_workbook,
)
from validation_checks.config import Settings
from validation_checks.main import build_pipeline, main
from validation_checks.models import CheckId, ExclusionScope, FileReport, FileStatus, MatchStatus
from validation_checks.pipeline import Pipeline
from validation_checks.report import write_report

SAMPLE = Path(__file__).parent.parent / "inputs" / "Argentina_Carrier_Survey_2027_Validation.xlsx"
PEERS = ["Carrier A", "Carrier B", "Carrier C", "Carrier D", "Carrier E", "Carrier F"]


@pytest.fixture
def pipeline(tmp_path: Path) -> Pipeline:
    return build_pipeline(SETTINGS.model_copy(update={"model_path": tmp_path / "no-model"}))


@pytest.mark.skipif(not SAMPLE.exists(), reason="sample workbook not present")
def test_argentina_sample(pipeline: Pipeline) -> None:
    report = pipeline.run(SAMPLE)

    assert {(e.sheet, e.column, e.scope) for e in report.exclusions} == {
        ("1.Country Peers", "Peer 4", ExclusionScope.CELL),
        ("2. Country LOB MAPPING.", "Property", ExclusionScope.CELL),
        ("2. Country LOB MAPPING.", "Surety", ExclusionScope.CELL),
    }
    assert [issue.check for issue in report.issues] == [CheckId.COUNTRY_MISMATCH]
    missing = {m.name.display for m in report.matches if m.hit.status is MatchStatus.NOT_FOUND}
    assert missing == {f"Carrier {letter}" for letter in "ABCEFG"}


def test_red_marks_flow_through_to_checks_and_matching(tmp_path: Path, pipeline: Pipeline) -> None:
    workbook = survey_workbook(
        peers=[peers_row("CHUBB", PEERS), peers_row("Excluded Co", ["Only one"])],
        lob=[lob_row("CHUBB", ["X", None, None, None]), *(lob_row(p, ["X"]) for p in PEERS)],
        panel=[panel_row(1, "a@x.com"), panel_row(2, "")],
    )
    workbook["1.Country Peers"]["D9"].fill = solid("FFC00000")  # dark red: row excluded
    workbook["1.Country Peers"]["G8"].font = red_font()  # Carrier A dropped, 5 peers remain
    workbook["3. User Panel"]["B9"].font = red_font()  # employee with blank email excluded

    report = pipeline.run(save(workbook, tmp_path / "Argentina_Carrier_Survey_2027.xlsx"))

    assert report.issues == ()
    assert len(report.exclusions) == 3
    assert [m.name.display for m in report.matches if m.needs_review] == ["Carrier A"]


def test_report_always_has_every_sheet(tmp_path: Path) -> None:
    reports = [FileReport.failed("bad.txt", "Not an Excel workbook")]

    path = write_report(reports, tmp_path, datetime(2026, 10, 14, 9, 30))

    workbook = load_workbook(path)
    assert path.name == "validation_results_20261014_093000.xlsx"
    assert workbook.sheetnames == ["Summary", "Issues", "LOB Matching", "Exclusions", "File Errors"]
    assert workbook["Summary"]["C2"].value == FileStatus.FAILED
    assert workbook["Issues"].max_row == 1


def test_cli_runs_every_input_file_and_logs(tmp_path: Path) -> None:
    inputs, output = tmp_path / "in", tmp_path / "out"
    inputs.mkdir()
    for country in ("Peru", "Chile", "Mexico"):
        workbook = survey_workbook(peers=[peers_row("CHUBB", PEERS)], lob=[], panel=[])
        save(workbook, inputs / f"{country}_Carrier_Survey_2027.xlsx")
    (inputs / "~$Peru_Carrier_Survey_2027.xlsx").write_bytes(b"lock file")
    (inputs / "notes.txt").write_text("not a survey", encoding="utf-8")
    args = ["--input", str(inputs), "--output", str(output), "--model-path", str(tmp_path / "m")]

    assert main(args) == 0

    (results,) = output.glob("validation_results_*.xlsx")
    summary = load_workbook(results)["Summary"]
    assert sorted(str(row[0]) for row in summary.iter_rows(min_row=2, values_only=True)) == [
        "Chile_Carrier_Survey_2027.xlsx",
        "Mexico_Carrier_Survey_2027.xlsx",
        "Peru_Carrier_Survey_2027.xlsx",
        "notes.txt",
    ]
    (log_file,) = output.glob("validation_run_*.log")
    log = log_file.read_text(encoding="utf-8")
    assert "Validating Peru_Carrier_Survey_2027.xlsx" in log
    assert "notes.txt: Not an Excel workbook" in log


def test_cli_without_input_files_exits_with_2(tmp_path: Path) -> None:
    assert main(["--input", str(tmp_path / "nope"), "--output", str(tmp_path)]) == 2


def test_settings_defaults_point_at_project_folders() -> None:
    defaults = Settings()

    assert (defaults.input_dir, defaults.output_dir, defaults.model_path) == (
        Path("inputs"),
        Path("output"),
        Path("sentence-transformers"),
    )
